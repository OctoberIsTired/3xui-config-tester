import json
import zipfile
from pathlib import Path

import pytest

from app.results.writer import ResultStore
from app.security.masking import mask_inbound_payload, mask_log_text, mask_parameter_values, mask_secrets
from app.parameters.models import ParameterSpec
from app.state.checkpoint import Checkpoint, atomic_json


def test_atomic_json_retries_replace_while_the_destination_is_open(tmp_path: Path, monkeypatch) -> None:
    """Поллинг статуса держит state.json открытым и срывает os.replace с WinError 5
    (Windows не передаёт FILE_SHARE_DELETE). Короткий ретрай должен абсорбировать
    гонку, а не убивать прогон; неустранимый отказ остаётся ошибкой."""
    import app.state.checkpoint as checkpoint_module

    path = tmp_path / "state.json"
    path.write_text('{"previous": true}', encoding="utf-8")
    original_replace, failures = checkpoint_module.os.replace, {"left": 2}

    def flaky_replace(source, destination, **kwargs):
        if failures["left"] > 0:
            failures["left"] -= 1
            raise PermissionError(5, "Отказано в доступе")
        return original_replace(source, destination, **kwargs)

    monkeypatch.setattr(checkpoint_module.os, "replace", flaky_replace)
    monkeypatch.setattr(checkpoint_module.time, "sleep", lambda _seconds: None)
    atomic_json(path, {"retry": "ok"})
    assert '"retry": "ok"' in path.read_text(encoding="utf-8")

    def dead_replace(source, destination, **kwargs):
        raise PermissionError(5, "Отказано в доступе")

    monkeypatch.setattr(checkpoint_module.os, "replace", dead_replace)
    with pytest.raises(PermissionError):
        atomic_json(path, {"retry": "fail"})
    assert '"retry": "ok"' in path.read_text(encoding="utf-8")


def test_checkpoint_roundtrip_and_dynamic_result_export(tmp_path: Path) -> None:
    state = Checkpoint("demo", {"xray": "v26.9.9"}, {"abc:1"}, 1)
    path = tmp_path / "state.json"
    state.save(path)
    loaded = Checkpoint.load(path)
    assert loaded and loaded.completed_keys == {"abc:1"} and loaded.failed == 1
    store = ResultStore(tmp_path)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "abc",
                  "configuration": {"network": "ws", "alpn": ["h2", "http/1.1"]},
                  "result": {"status": "OK", "latency_avg_ms": 12}})
    store.export(["csv", "json", "xlsx"])
    assert "parameter.network" in (tmp_path / "results.csv").read_text(encoding="utf-8")
    assert (tmp_path / "summary.csv").exists() and (tmp_path / "results.json").exists() and (tmp_path / "results.xlsx").exists()
    from openpyxl import load_workbook
    sheet = load_workbook(tmp_path / "results.xlsx", read_only=False)["results"]
    rows = sheet.iter_rows()
    headers = [cell.value for cell in next(rows)]
    row = [cell.value for cell in next(rows)]
    assert row[headers.index("parameter.alpn")] == '["h2","http/1.1"]'
    assert sheet.freeze_panes == "G2"


def test_secret_masking() -> None:
    assert mask_secrets({"api_token": "very-secret", "nested": {"password": "x"}}) == {"api_token": "***", "nested": {"password": "***"}}
    spec = ParameterSpec.from_dict("custom", {"type": "string", "values": ["sensitive"],
                                              "path": ["streamSettings", "realitySettings", "privateKey"]})
    assert mask_parameter_values({"custom": "sensitive"}, [spec]) == {"custom": "***"}


def test_log_masking_hides_identity_but_keeps_the_reason() -> None:
    text = ("[Warning] failed to read PROXY protocol header for "
            "8c1b0e2a-1f4d-4f0e-9b7a-6d5c4b3a2f10 user=alice@example.com password=hunter2")
    masked = mask_log_text(text)
    assert "failed to read PROXY protocol header" in masked
    assert "8c1b0e2a" not in masked and "alice@" not in masked and "hunter2" not in masked


def test_skipped_candidate_is_not_counted_as_failed_measurement(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "skip", "configuration": {},
                  "result": {"status": "SKIPPED", "reason_code": "reality_target_or_sni_missing"}})
    store.append({"test_id": 2, "run": 1, "configuration_hash": "ok", "configuration": {},
                  "result": {"status": "OK", "score": 1.0}})
    summaries = {item["configuration_hash"]: item for item in store._summaries(store.records())}
    assert summaries["skip"]["skipped_tests"] == 1
    assert summaries["skip"]["failed_tests"] == 0
    assert summaries["ok"]["success_rate"] == 1.0


def test_xlsx_only_export_does_not_create_csv_and_accepts_structured_values(tmp_path: Path) -> None:
    directory = tmp_path / "xlsx-only"
    store = ResultStore(directory)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "abc",
                  "configuration": {"dest_override": ["http", "tls", "quic", "fakedns"]},
                  "result": {"status": "OK", "latency_avg_ms": 10.25}})
    store.export(["xlsx"])
    assert (directory / "results.xlsx").exists()
    assert not (directory / "results.csv").exists()
    assert not (directory / "summary.csv").exists()
    assert not (directory / "results.json").exists()


def test_xlsx_flattens_nested_metric_groups_into_columns(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "abc", "configuration": {},
                  "result": {"status": "FAILED", "stage": "screening",
                             "screening": {"success_rate": 0.0, "tcp_connect_ok": True},
                             "gate_failures": ["success_rate<1"]}})
    store.export(["xlsx"])
    from openpyxl import load_workbook
    sheet = load_workbook(tmp_path / "results.xlsx")["results"]
    headers = [cell.value for cell in sheet[1]]
    assert "screening.success_rate" in headers
    assert "screening.tcp_connect_ok" in headers
    assert sheet.cell(2, headers.index("gate_failures") + 1).value == '["success_rate<1"]'


def test_summary_and_xlsx_dashboard_report_repeat_stability(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    for run, p95, speed in ((1, 100.0, 120.0), (2, 120.0, 140.0)):
        store.append({"test_id": 1, "run": run, "configuration_hash": "abc",
                      "configuration": {"tls_alpn": ["h2"]},
                      "result": {"status": "OK", "score": 0.9, "latency_p95_ms": p95,
                                 "download_mbps": speed}})
    summary = store._summaries(store.records())
    assert summary[0]["latency_p95_stddev"] > 0
    assert summary[0]["download_stddev"] > 0
    store.export(["xlsx"])
    from openpyxl import load_workbook
    workbook = load_workbook(tmp_path / "results.xlsx")
    assert workbook.sheetnames[:2] == ["Dashboard", "Candidates"]
    assert len(workbook["Dashboard"]._charts) == 4
    candidates = workbook["Candidates"]
    assert candidates["A2"].value == "#1"
    assert candidates["B2"].value == '{"tls_alpn": ["h2"]}'
    assert candidates["C2"].value == 1
    assert candidates["E2"].value == "abc"
    assert candidates["M2"].value > 0
    assert workbook["Dashboard"]["B43"].value == "#1"
    assert workbook["Dashboard"]["C43"].value == candidates["B2"].value
    score, latency, scatter, stability = workbook["Dashboard"]._charts
    assert score.series[0].val.numRef.f == "'Candidates'!$J$2"
    assert latency.series[0].val.numRef.f == "'Candidates'!$L$2"
    assert stability.series[0].val.numRef.f == "'Candidates'!$M$2"
    assert scatter.series[0].xVal.numRef.f == "'Candidates'!$L$2"
    assert scatter.series[0].yVal.numRef.f == "'Candidates'!$N$2"


def test_xlsx_chart_numbers_map_to_complete_configurations(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    for test_id, network, latency, speed in ((7, "ws", 95, 120), (8, "tcp", 110, 90)):
        store.append({"test_id": test_id, "run": 1, "configuration_hash": network,
                      "configuration": {"network": network, "tls_alpn": ["h2"]},
                      "result": {"status": "OK", "score": 0.8,
                                 "latency_p95_ms": latency, "download_mbps": speed}})
    store.export(["xlsx"])
    from openpyxl import load_workbook

    workbook = load_workbook(tmp_path / "results.xlsx")
    rows = list(workbook["Candidates"].values)
    assert [row[0] for row in rows[1:]] == ["#1", "#2"]
    assert [row[2] for row in rows[1:]] == [7, 8]
    assert '"network": "ws"' in rows[1][1]
    assert '"network": "tcp"' in rows[2][1]
    assert [workbook["Dashboard"].cell(row, 2).value for row in (43, 44)] == ["#1", "#2"]
    scatter = workbook["Dashboard"]._charts[2]
    assert len(scatter.series) == 2
    assert [series.title.strRef.strCache.pt[0].v if series.title.strRef else series.title.v
            for series in scatter.series] == ["#1", "#2"]


def test_exported_candidate_rows_follow_journal_order(tmp_path: Path) -> None:
    store = ResultStore(tmp_path)
    for test_id, score, phase in ((1, 0.2, "search"), (2, 0.9, "search"), (3, 0.5, "validation")):
        store.append({"test_id": test_id, "run": 1, "phase": phase,
                      "configuration_hash": str(test_id), "configuration": {"candidate": test_id},
                      "result": {"status": "OK", "score": score}})

    assert [item["test_id"] for item in store._summaries(store.records())] == [1, 2, 3]
    store.export(["xlsx"])
    from openpyxl import load_workbook
    workbook = load_workbook(tmp_path / "results.xlsx")
    assert [row[2] for row in list(workbook["Candidates"].values)[1:]] == [1, 2, 3]
    assert [workbook["Dashboard"].cell(row, 2).value for row in (43, 44, 45)] == ["#1", "#2", "#3"]


def test_xlsx_charts_drop_per_point_data_labels_when_labels_overlap(tmp_path: Path) -> None:
    """Подпись у каждой точки накладывается на соседнюю у диаграмм, которые
    читают по осям; подписи снимаются, а подписи осей и заголовки остаются.
    Scatter подписывает точки их номерами — там подпись это тождество, не значение."""
    store = ResultStore(tmp_path)
    for run, p95, speed in ((1, 100.0, 120.0), (2, 120.0, 140.0)):
        store.append({"test_id": 1, "run": run, "configuration_hash": "abc",
                      "configuration": {"tls_alpn": ["h2"]},
                      "result": {"status": "OK", "score": 0.9, "latency_p95_ms": p95,
                                 "download_mbps": speed}})
    store.append({"test_id": 2, "run": 1, "configuration_hash": "def",
                  "configuration": {"network": "ws"},
                  "result": {"status": "OK", "score": 0.5, "latency_p95_ms": 200.0, "download_mbps": 50.0}})
    store.export(["xlsx"])
    from openpyxl import load_workbook

    workbook = load_workbook(tmp_path / "results.xlsx")
    score, latency, scatter, stability = workbook["Dashboard"]._charts
    for chart in (score, latency, scatter, stability):
        assert all(series.dLbls is None for series in chart.series)
    assert score.dLbls is None and latency.dLbls is None and stability.dLbls is None
    assert score.title is not None and score.x_axis.title is not None and score.y_axis.title is not None
    assert scatter.dLbls is not None and scatter.dLbls.showSerName is True

    parts = {}
    with zipfile.ZipFile(tmp_path / "results.xlsx") as archive:
        for name in archive.namelist():
            if name.startswith("xl/charts/"):
                parts[name] = archive.read(name).decode("utf-8")
    assert len(parts) == 4
    for title in ("Средний балл по конфигурациям", "Средняя задержка p95", "Разброс задержки p95"):
        xml = next(value for value in parts.values() if f"<a:t>{title}</a:t>" in value)
        assert "dLbls" not in xml and "showVal" not in xml
    scatter_xml = next(value for value in parts.values() if "<a:t>Скорость и задержка p95</a:t>" in value)
    assert "showSerName" in scatter_xml


def _inbound_payload(port: int = 8443) -> dict:
    return mask_inbound_payload({
        "id": 9, "up": 0, "down": 0, "expiryTime": 0,
        "remark": "[3xui-config-tester:0f1f1f1f1f1f1f1f1f1f1f1f1f1f1f1f]",
        "enable": True, "listen": "", "port": port, "protocol": "vless",
        "settings": json.dumps({"clients": [{"id": "8c1b0e2a-1f4d-4f0e-9b7a-6d5c4b3a2f10", "flow": "",
                                             "email": "3xui-tester-aa@example.invalid", "enable": True}],
                                "decryption": "none"}, separators=(",", ":")),
        "streamSettings": json.dumps({"network": "tcp", "security": "tls",
                                      "tlsSettings": {"serverName": "example.test",
                                                      "certificates": [{"keyFile": "/etc/ssl/panel.key"}]}},
                                     separators=(",", ":")),
        "sniffing": json.dumps({"enabled": True, "destOverride": ["http", "tls"]}, separators=(",", ":")),
    })


def test_inbound_payload_template_masks_credentials_and_keeps_structure() -> None:
    """Шаблон inbound'а попадает в артефакты, поэтому учётные данные в нём
    заменены, а структура (порт, протокол, транспорт, sockopt) сохранена."""
    payload = dict(_inbound_payload())
    payload["streamSettings"] = json.dumps({"network": "kcp", "security": "none",
                                            "kcpSettings": {"seed": "private-seed", "header": {"type": "utp"}},
                                            "sockopt": {"domainStrategy": "AsIs"}}, separators=(",", ":"))
    template = mask_inbound_payload(payload)
    assert "id" not in template and "up" not in template and "expiryTime" not in template
    assert template["port"] == 8443 and template["protocol"] == "vless" and template["enable"] is True
    client = json.loads(template["settings"])["clients"][0]
    assert client["id"] == "***" and client["email"] == "***"
    kcp = json.loads(template["streamSettings"])["kcpSettings"]
    assert kcp["seed"] == "***" and kcp["header"] == {"type": "utp"}
    assert json.loads(template["streamSettings"])["sockopt"] == {"domainStrategy": "AsIs"}

    payload["streamSettings"] = json.dumps({"network": "tcp", "security": "reality",
                                            "realitySettings": {"dest": "example.test:443",
                                                                "privateKey": "PRIVATE-KEY",
                                                                "shortIds": ["deadbeef", "cafebabe"],
                                                                "settings": {"publicKey": "PUBLIC-KEY"}},
                                            "sockopt": {"domainStrategy": "AsIs"}}, separators=(",", ":"))
    template = mask_inbound_payload(payload)
    reality = json.loads(template["streamSettings"])["realitySettings"]
    assert reality["dest"] == "example.test:443" and reality["privateKey"] == "***"
    assert reality["shortIds"] == ["***", "***"] and reality["settings"]["publicKey"] == "***"
    assert json.loads(template["streamSettings"])["sockopt"] == {"domainStrategy": "AsIs"}
    assert "8c1b0e2a" not in json.dumps(template, ensure_ascii=False, sort_keys=True)
    assert json.dumps(template, ensure_ascii=False)  # остаётся JSON-безопасным


def test_xlsx_inbound_sheet_offers_one_copy_pasteable_inbound_per_candidate(tmp_path: Path) -> None:
    """Полезную нагрузку можно скопировать в редактор inbound'а 3x-ui, поэтому
    она лежит отдельным листом: по строке на конфигурацию, без повторов повторов."""
    store = ResultStore(tmp_path)
    payload = _inbound_payload()
    for run in (1, 2):
        store.append({"test_id": 1, "run": run, "configuration_hash": "h1",
                      "configuration": {"tls_alpn": ["h2"]}, "inbound_payload": payload,
                      "result": {"status": "OK", "score": 0.9, "latency_p95_ms": 100.0, "download_mbps": 10.0}})
    store.append({"test_id": 2, "run": 1, "configuration_hash": "h2",
                  "configuration": {"network": "ws"}, "inbound_payload": {**payload, "port": 24443},
                  "result": {"status": "OK", "score": 0.5, "latency_p95_ms": 200.0, "download_mbps": 5.0}})
    store.export(["csv", "json", "xlsx"])
    from openpyxl import load_workbook

    workbook = load_workbook(tmp_path / "results.xlsx")
    assert workbook.sheetnames[:3] == ["Dashboard", "Candidates", "Inbound"]
    sheet = workbook["Inbound"]
    headers = [cell.value for cell in sheet[1]]
    column = next(index for index, header in enumerate(headers) if header.startswith("Inbound JSON"))
    rows = list(sheet.values)[1:]
    assert [row[headers.index("configuration_hash")] for row in rows] == ["h1", "h2"]
    assert [row[0] for row in rows] == ["#1", "#2"]

    template = json.loads(str(rows[0][column]))
    assert template["port"] == 8443 and template["protocol"] == "vless"
    assert json.loads(str(rows[0][headers.index("settings (JSON)")]))["clients"][0]["id"] == "***"
    assert rows[0][headers.index("configuration_hash")] == "h1" and "8c1b0e2a" not in str(rows[0][column])
    # payload не дублируется в построчных листах и CSV, где он не нужен.
    assert "inbound_payload" not in [cell.value for cell in workbook["results"][1]]
    assert "inbound_payload" not in (tmp_path / "results.csv").read_text(encoding="utf-8")
    assert json.loads((tmp_path / "results.json").read_text(encoding="utf-8"))[0]["inbound_payload"]["port"] == 8443


def test_xlsx_without_inbound_payload_keeps_the_previous_sheets(tmp_path: Path) -> None:
    """Записи старых прогонов payload'а не несут: лист не появляется, и порядок
    листов прежнего экспорта не меняется."""
    store = ResultStore(tmp_path)
    store.append({"test_id": 1, "run": 1, "configuration_hash": "abc", "configuration": {},
                  "result": {"status": "OK", "score": 1.0}})
    store.export(["xlsx"])
    from openpyxl import load_workbook

    workbook = load_workbook(tmp_path / "results.xlsx")
    assert workbook.sheetnames[:2] == ["Dashboard", "Candidates"]
    assert "Inbound" not in workbook.sheetnames
